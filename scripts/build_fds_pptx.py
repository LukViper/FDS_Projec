#!/usr/bin/env python3
"""Build FDS.pptx — Foundations of Data Science project presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parents[1] / "FDS.pptx"

# Palette — clean academic, not purple-AI default
NAVY = RGBColor(0x1B, 0x3A, 0x4B)
TEAL = RGBColor(0x2A, 0x6F, 0x7F)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
LIGHT = RGBColor(0xF7, 0xF4, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x2C, 0x2C, 0x2C)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
ROW_ALT = RGBColor(0xEE, 0xF3, 0xF5)

W, H = Inches(13.333), Inches(7.5)  # widescreen 16:9


def set_run(run, *, size=18, bold=False, color=DARK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bg(slide, color=LIGHT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_top_bar(slide, title: str, subtitle: str | None = None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.18), Inches(12.3), Inches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, size=24, bold=True, color=WHITE, font="Calibri")

    if subtitle:
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(0.85), W, Inches(0.38)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL
        accent.line.fill.background()
        sbox = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.88), Inches(12.3), Inches(0.32)
        )
        stf = sbox.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle
        set_run(sr, size=14, color=WHITE)


def add_footer(slide, page: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(10), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Foundations of Data Science  ·  Customer Churn Prediction"
    set_run(r, size=11, color=MUTED)

    num = slide.shapes.add_textbox(Inches(11.5), Inches(7.1), Inches(1.5), Inches(0.3))
    ntf = num.text_frame
    ntf.clear()
    np = ntf.paragraphs[0]
    np.alignment = PP_ALIGN.RIGHT
    nr = np.add_run()
    nr.text = f"{page} / {total}"
    set_run(nr, size=11, color=MUTED)


def bullets(slide, items, *, left=0.5, top=1.5, width=12.3, height=5.2, size=18, spacing=8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(spacing)
        if isinstance(item, tuple):
            text, level = item
            p.level = level
        else:
            text = item
        # support bold prefix via **text**
        if text.startswith("**") and "**" in text[2:]:
            end = text.index("**", 2)
            bold_part = text[2:end]
            rest = text[end + 2 :]
            r1 = p.add_run()
            r1.text = "•  " + bold_part
            set_run(r1, size=size, bold=True, color=DARK)
            if rest:
                r2 = p.add_run()
                r2.text = rest
                set_run(r2, size=size, color=DARK)
        else:
            r = p.add_run()
            r.text = "•  " + text
            set_run(r, size=size, color=DARK)
    return box


def add_table(slide, rows, col_widths, *, left=0.5, top=1.5, height=None):
    n_rows, n_cols = len(rows), len(rows[0])
    total_w = sum(col_widths)
    tbl_h = height or min(0.42 * n_rows + 0.1, 5.2)
    shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(left), Inches(top), Inches(total_w), Inches(tbl_h)
    )
    table = shape.table
    for j, w in enumerate(col_widths):
        table.columns[j].width = Inches(w)

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(cell_text)
            if i == 0:
                set_run(run, size=13, bold=True, color=WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                set_run(run, size=12, color=DARK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if i % 2 else ROW_ALT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape


def notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    slides_meta = []  # (builder fn collecting into list later — we'll count)

    # ---------- Slide builders return notes text ----------
    builders = []

    def title_slide():
        s = blank(prs)
        add_bg(s, NAVY)
        # accent stripe
        stripe = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(5.6), W, Inches(1.9)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = TEAL
        stripe.line.fill.background()

        t = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(12), Inches(1.2))
        tf = t.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = "Customer Churn Prediction"
        set_run(r, size=40, bold=True, color=WHITE, font="Calibri")

        t2 = s.shapes.add_textbox(Inches(0.7), Inches(2.9), Inches(12), Inches(0.7))
        p2 = t2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = "with Explainable AI"
        set_run(r2, size=28, color=RGBColor(0xD4, 0xE8, 0xEC), font="Calibri")

        t3 = s.shapes.add_textbox(Inches(0.7), Inches(4.0), Inches(12), Inches(0.5))
        p3 = t3.text_frame.paragraphs[0]
        r3 = p3.add_run()
        r3.text = "Foundations of Data Science  ·  CO1–CO5"
        set_run(r3, size=20, bold=True, color=ACCENT, font="Calibri")

        t4 = s.shapes.add_textbox(Inches(0.7), Inches(5.85), Inches(12), Inches(1.2))
        tf4 = t4.text_frame
        p4 = tf4.paragraphs[0]
        r4 = p4.add_run()
        r4.text = "UCI Online Retail II  ·  4-member team project"
        set_run(r4, size=18, color=WHITE)
        p5 = tf4.add_paragraph()
        r5 = p5.add_run()
        r5.text = "Pipeline: Preprocessing → EDA & Features → Classification → Certainty & Demo"
        set_run(r5, size=15, color=RGBColor(0xD4, 0xE8, 0xEC))

        notes(
            s,
            "Good morning. Today we walk through our Foundations of Data Science project "
            "— not as an ML demo, but as a full workflow: messy transactions → clean labels → "
            "exploration → careful modelling → how sure we can be about what we found.",
        )
        return s

    def content_slide(title, subtitle, body_fn, note_text, page, total):
        s = blank(prs)
        add_bg(s)
        add_top_bar(s, title, subtitle)
        body_fn(s)
        add_footer(s, page, total)
        notes(s, note_text)
        return s

    # Build slide list conceptually then create with page numbers
    # We'll create all slides in order; total known in advance.
    TOTAL = 25

    # 1 Title
    title_slide()

    # 2 Why
    def s2(s):
        bullets(
            s,
            [
                "Online retailers lose revenue when customers stop buying",
                "Our question: Among customers active in 2010, who does not return in early 2011?",
                "Goal: a trustworthy, time-aware churn analysis — then predict and explain",
                "Most of the work is Foundations of Data Science: make the problem valid before models",
            ],
            top=1.5,
            size=20,
            spacing=14,
        )

    content_slide(
        "Why this project?",
        "Problem in plain words",
        s2,
        "Business problem is simple; data-science challenge is harder. Raw file is invoice lines, "
        "not churn yes/no. Turning events into a valid customer-level problem is what FDS trains us to do.",
        2,
        TOTAL,
    )

    # 3 Organisation
    def s3(s):
        bullets(
            s,
            [
                "M1 Data & Preprocessing → M2 EDA & Features → M3 Classification → M4 Certainty & Demo",
            ],
            top=1.4,
            size=16,
            spacing=6,
            height=0.6,
        )
        add_table(
            s,
            [
                ["Member", "Owns", "CO weight"],
                ["Member 1 — Shasank", "Cleaning, periods, churn labels", "CO1, CO2"],
                ["Member 2 — [Name]", "EDA, visuals, RFM / behaviour features", "CO2, CO3"],
                ["Member 3 — [Name]", "Temporal split, LR / RF / XGBoost", "CO4"],
                ["Member 4 — [Name]", "Statistical tests, SHAP, Streamlit", "CO1, CO5"],
            ],
            [3.5, 5.5, 2.5],
            top=2.2,
        )

    content_slide(
        "How we organised the work",
        "Clear ownership · sequential handoffs",
        s3,
        "Each person owns a real stage. Without Member 1's labels, nothing else is valid. "
        "Faculty can verify each stage with files and numbers.",
        3,
        TOTAL,
    )

    # 4 Dataset
    def s4(s):
        add_table(
            s,
            [
                ["Item", "Detail"],
                ["Source", "UCI Online Retail II (Chen, 2019)"],
                ["Raw size", "~1.07 million invoice lines"],
                ["Columns", "Invoice, StockCode, Quantity, InvoiceDate, Price, Customer ID, Country…"],
                ["Span", "Dec 2009 → Dec 2011"],
                ["Study design", "Observe 2010 · Label repurchase Jan–Jun 2011"],
            ],
            [3.0, 9.5],
            top=1.5,
        )

    content_slide(
        "Dataset snapshot",
        "Common to all members",
        s4,
        "We started from a real public retail dataset — identity gaps, cancellations, "
        "duplicates, and time windows show up in practice.",
        4,
        TOTAL,
    )

    # 5 CO map
    def s5(s):
        add_table(
            s,
            [
                ["CO", "Meaning", "In our project"],
                ["CO1", "Statistical foundations", "Binary outcome; distributions; significance; certainty language"],
                ["CO2", "Pre-processing raw data", "IDs / cancels / dupes; imputation; observation-only features"],
                ["CO3", "EDA & insightful visuals", "Churn mix, country, RFM plots, correlation with churn"],
                ["CO4", "ML for classification / insight", "Temporal split; LR vs RF vs XGBoost; metrics beyond accuracy"],
                ["CO5", "Degree of certainty", "Mann–Whitney, point-biserial, chi-square; SHAP explanations"],
            ],
            [1.2, 3.5, 7.5],
            top=1.4,
        )

    content_slide(
        "Course outcomes map",
        "Our compass for this presentation",
        s5,
        "We follow COs in order. ML appears only after statistics, cleaning, and exploration — "
        "models are one way to get insight, not the whole subject.",
        5,
        TOTAL,
    )

    # 6 M1 intro
    def s6(s):
        bullets(
            s,
            [
                "**Owned by Member 1 — Shasank**",
                "Acquire and cite the UCI dataset",
                "Inspect structure and data types",
                "Define cleaning rules and document why",
                "Define observation vs prediction periods",
                "Define churn scientifically",
                "Hand Member 2 a validated customer table",
                "Artifacts: src/preprocessing.py · data/processed/ · docs/preprocessing_decisions.md",
            ],
            top=1.45,
            size=18,
            spacing=10,
        )

    content_slide(
        "Member 1 — Shasank",
        "Data Collection & Preprocessing  ·  CO1 + CO2",
        s6,
        "My job was to make the problem valid. If churn labels are wrong, every chart and model later is wrong.",
        6,
        TOTAL,
    )

    # 7 Cleaning
    def s7(s):
        add_table(
            s,
            [
                ["Cleaning step", "Rows affected", "Why"],
                ["Missing Customer ID", "243,007", "Guests cannot be tracked → cannot define churn"],
                ["Cancelled invoices (C…)", "18,744", "Cancellations are not completed purchases"],
                ["Invalid qty / price ≤ 0", "71", "Not reliable sales"],
                ["Exact duplicates", "26,124", "Recording noise, not extra purchases"],
                ["Cleaned transactions", "779,425", "Usable purchase events"],
            ],
            [3.8, 2.2, 6.2],
            top=1.45,
        )

    content_slide(
        "From invoices to customers",
        "CO2 in practice — every drop has a reason",
        s7,
        "Pre-processing is not deleting inconvenient rows. Each rule has a definitional reason. "
        "Without Customer ID we cannot follow the same person into 2011.",
        7,
        TOTAL,
    )

    # 8 Churn design
    def s8(s):
        bullets(
            s,
            [
                "Observation: 1 Jan 2010 — 31 Dec 2010  → behaviour we may use as inputs",
                "Prediction:  1 Jan 2011 — 30 Jun 2011  → used ONLY for the label",
                "Eligible: bought at least once in 2010",
                "churn = 1 if no cleaned purchase in Jan–Jun 2011; else 0",
                "Result: 4,231 customers · 2,217 churned · 2,014 retained · ≈ 52.4% churn",
                "This separation avoids target leakage — CO1 study design before modelling",
            ],
            top=1.45,
            size=18,
            spacing=12,
        )

    content_slide(
        "Study design: time windows & churn",
        "CO1 — statistical foundations of the outcome",
        s8,
        "Churn is a carefully defined binary outcome. Past window for features, future window for labels.",
        8,
        TOTAL,
    )

    # 9 M1 learnings
    def s9(s):
        bullets(
            s,
            [
                "Real data is messy; cleaning rules must be written down",
                "Labels are a research-design decision — not a free column",
                "Validation matters: M1 report passed all_passed = true before handoff",
                "Practical FDS: good data science starts with definitions and honesty about drops",
            ],
            top=1.6,
            size=20,
            spacing=16,
        )

    content_slide(
        "What Member 1 learned",
        "Practical Foundations of Data Science",
        s9,
        "Six-month prediction window: long enough to see repurchase, still inside available calendar.",
        9,
        TOTAL,
    )

    # 10 M2 intro
    def s10(s):
        bullets(
            s,
            [
                "**Owned by Member 2 — [Name]**",
                "Validate Member 1’s handoff (churn labels unchanged)",
                "Explore customer and churn distributions",
                "Build RFM + behavioural features from 2010 only",
                "Handle feature missingness carefully",
                "Produce visuals and feature dictionary for Member 3",
                "Artifacts: src/features.py · data/features/ · docs/m2_eda_findings.md",
            ],
            top=1.45,
            size=18,
            spacing=10,
        )

    content_slide(
        "Member 2 — [Name]",
        "EDA & Feature Engineering  ·  CO2 + CO3",
        s10,
        "I take Member 1’s customer list and ask what these people look like — and how to summarise a year of shopping into readable numbers.",
        10,
        TOTAL,
    )

    # 11 EDA
    def s11(s):
        bullets(
            s,
            [
                "Near-balanced churn: ~52% vs ~48%",
                "UK dominates country mix (visible in top-10 chart)",
                "Monetary & frequency are right-skewed — typical retail, not “bad data”",
            ],
            top=1.35,
            size=17,
            spacing=8,
            height=1.8,
        )
        add_table(
            s,
            [
                ["Feature", "Corr. with churn", "Story"],
                ["Recency (days since last buy)", "+0.31", "Longer silence → more churn"],
                ["Product diversity", "−0.29", "Broader basket → less churn"],
                ["Frequency", "−0.27", "More invoices → less churn"],
            ],
            [4.5, 2.5, 5.2],
            top=3.4,
        )

    content_slide(
        "Exploratory findings",
        "CO3 — patterns before algorithms",
        s11,
        "Before any algorithm, pictures already tell a story. EDA guides features and later hypothesis tests.",
        11,
        TOTAL,
    )

    # 12 Features
    def s12(s):
        bullets(
            s,
            [
                "**RFM:** Recency · Frequency · Monetary (relative to 31 Dec 2010)",
                "**Behavioural:** AOV, purchase interval, product diversity, cancellation rate, spending/order trends",
                "Hard rule: no prediction-period fields · churn_unchanged = true",
                "Single-purchase interval: median impute + flag (~1,418 customers)",
                "IQR outliers reported but kept — do not distort churn base rate",
                "Cancellation rate from raw C-invoices (cleaned file has none)",
            ],
            top=1.45,
            size=18,
            spacing=12,
        )

    content_slide(
        "Feature engineering as pre-processing",
        "CO2 continued — analysis-ready variables without leakage",
        s12,
        "Feature engineering is still Foundations work: transform events, defend imputation, control leakage.",
        12,
        TOTAL,
    )

    # 13 M2 learnings
    def s13(s):
        bullets(
            s,
            [
                "EDA first saves modelling from blind guessing",
                "Correlation is a clue, not causation",
                "Observation-only features are the line between science and cheating the future",
                "A feature dictionary matters as much as code",
            ],
            top=1.6,
            size=20,
            spacing=16,
        )

    content_slide(
        "What Member 2 learned",
        "Practical Foundations of Data Science",
        s13,
        "If Member 3 cannot explain what recency_days means, the model is already a black box.",
        13,
        TOTAL,
    )

    # 14 M3 intro
    def s14(s):
        bullets(
            s,
            [
                "**Owned by Member 3 — [Name]**  (CO4 — kept secondary to FDS)",
                "Temporal train / validation / test split",
                "Class-imbalance awareness under time shift",
                "Train Logistic Regression, Random Forest, XGBoost",
                "Compare with ROC-AUC, PR-AUC, Precision, Recall, F1 — not accuracy alone",
                "Freeze selected pipeline for Member 4",
                "Artifacts: src/modeling.py · models/ · docs/experiment_notes.md",
            ],
            top=1.45,
            size=18,
            spacing=10,
        )

    content_slide(
        "Member 3 — [Name]",
        "Classification for insight  ·  CO4",
        s14,
        "Important part is how we evaluate under time, not who wins a contest.",
        14,
        TOTAL,
    )

    # 15 Temporal
    def s15(s):
        add_table(
            s,
            [
                ["Split", "Rule (last purchase in 2010)", "Role"],
                ["Train", "≤ 31 Aug 2010", "Fit preprocess + model"],
                ["Val", "Sep–Oct 2010", "Select model"],
                ["Test", "≥ 1 Nov 2010", "Honest final check"],
            ],
            [2.5, 5.5, 4.2],
            top=1.4,
        )
        bullets(
            s,
            [
                "Churn rate shifts across cohorts (train ~74% → test ~37%)",
                "Used class weights / scale_pos_weight — world changes over time",
                "Random shuffle would mix cohorts and give optimistic numbers",
            ],
            top=4.0,
            size=17,
            spacing=10,
            height=2.5,
        )

    content_slide(
        "Temporal protocol & imbalance",
        "Still FDS thinking — honest evaluation",
        s15,
        "Temporal splitting is how Foundations of Data Science meets real sequential data.",
        15,
        TOTAL,
    )

    # 16 Results
    def s16(s):
        add_table(
            s,
            [
                ["Model", "Val PR-AUC", "Test ROC-AUC", "Selected?"],
                ["Logistic Regression", "0.691", "0.715", "No"],
                ["Random Forest", "0.725", "0.755", "Yes"],
                ["XGBoost", "0.691", "0.709", "No"],
            ],
            [4.0, 2.5, 3.0, 2.5],
            top=1.4,
        )
        bullets(
            s,
            [
                "Selected RF test: PR-AUC 0.625 · Precision 0.646 · Recall 0.494 · F1 0.560",
                "We use these numbers to ask: do behavioural features carry signal? — not to sell magic",
            ],
            top=4.0,
            size=17,
            spacing=12,
            height=2.2,
        )

    content_slide(
        "Model comparison",
        "CO4 — insight, briefly",
        s16,
        "Random Forest won on validation PR-AUC. CO4 in service of insight, not the centre of the course.",
        16,
        TOTAL,
    )

    # 17 M3 learnings
    def s17(s):
        bullets(
            s,
            [
                "Algorithm choice matters less than split design and metric choice",
                "Accuracy alone misleads under drifting base rates",
                "Pipelines must freeze preprocessing with the model",
                "In FDS, modelling is the last mile of a careful experiment — not the first slide",
            ],
            top=1.6,
            size=20,
            spacing=16,
        )

    content_slide(
        "What Member 3 learned",
        "Practical Foundations of Data Science",
        s17,
        "Modelling is the last mile of a careful experiment.",
        17,
        TOTAL,
    )

    # 18 M4 intro
    def s18(s):
        bullets(
            s,
            [
                "**Owned by Member 4 — [Name]**",
                "Load Member 3’s frozen Random Forest (no silent retrain)",
                "Hypothesis tests: feature ↔ churn association",
                "Explain predictions with SHAP (global + local)",
                "Streamlit demo: probability, risk band, top factors",
                "Artifacts: src/explain.py · app/streamlit_app.py · docs/xai/",
            ],
            top=1.45,
            size=18,
            spacing=12,
        )

    content_slide(
        "Member 4 — [Name]",
        "Certainty, explanation & delivery  ·  CO1 + CO5",
        s18,
        "CO5: analyse degree of certainty using statistical tests and models — same labels, features, frozen classifier.",
        18,
        TOTAL,
    )

    # 19 Stats
    def s19(s):
        add_table(
            s,
            [
                ["Test", "Applied to", "α"],
                ["Mann–Whitney U + point-biserial", "12 numeric features vs churn", "0.05"],
                ["Chi-square", "Country (top groups + Other)", "0.05"],
            ],
            [5.0, 5.0, 2.0],
            top=1.4,
        )
        bullets(
            s,
            [
                "Result: 12/12 numeric features significant; country also significant",
                "Tests = feature differs by churn label in the data",
                "SHAP = what the trained model uses",
                "Together: data-level association + model-level attribution — not the same claim",
            ],
            top=3.5,
            size=17,
            spacing=10,
            height=3.0,
        )

    content_slide(
        "Statistical certainty",
        "CO5 / CO1 — how sure are we?",
        s19,
        "We do not say the model is right because AUC looks nice. We ask whether engineered features "
        "actually separate churned and retained customers under nonparametric tests.",
        19,
        TOTAL,
    )

    # 20 SHAP
    def s20(s):
        bullets(
            s,
            [
                "Global drivers: monetary, frequency, average order value, spending trend…",
                "Local: for one high-risk customer, show factors pushing probability up or down",
                "Streamlit risk bands: Low < 0.33 · Medium · High ≥ 0.66",
                "In an FDS talk: explainability supports trust; tests support certainty language",
            ],
            top=1.6,
            size=20,
            spacing=16,
        )

    content_slide(
        "SHAP in one minute",
        "Supporting insight for stakeholders",
        s20,
        "SHAP helps a stakeholder hear why this customer looks risky.",
        20,
        TOTAL,
    )

    # 21 M4 learnings
    def s21(s):
        bullets(
            s,
            [
                "“Significant” and “important to the model” are related but not identical",
                "Freezing the pipeline prevents explanation theatre",
                "A demo is useful only if inputs still respect the observation window",
                "Certainty means knowing — and showing — the limits of your claims",
            ],
            top=1.6,
            size=20,
            spacing=16,
        )

    content_slide(
        "What Member 4 learned",
        "Practical Foundations of Data Science",
        s21,
        "Certainty means knowing the limits of your claims — and showing them.",
        21,
        TOTAL,
    )

    # 22 End-to-end
    def s22(s):
        bullets(
            s,
            [
                "1.07M raw lines",
                "→ 779k cleaned purchases",
                "→ 4,231 labelled customers (~52% churn)",
                "→ RFM + behaviour features (observation-only)",
                "→ Temporally evaluated RF (test ROC-AUC ~0.76)",
                "→ Significant feature–churn tests + SHAP + Streamlit",
                "That pipeline is the project. Models are one station on the line.",
            ],
            top=1.45,
            size=19,
            spacing=12,
        )

    content_slide(
        "End-to-end result",
        "One picture of the whole journey",
        s22,
        "That pipeline is the project. Models are one station on the line.",
        22,
        TOTAL,
    )

    # 23 CO learnings
    def s23(s):
        add_table(
            s,
            [
                ["CO", "Practical lesson from this project"],
                ["CO1", "Study design: windows, binary outcomes, distributions, significance, careful certainty wording"],
                ["CO2", "Pre-processing = reasoning under constraints — documented so others can reproduce"],
                ["CO3", "EDA surfaces patterns (recency, frequency, diversity) before algorithms"],
                ["CO4", "Classification helps with temporal evaluation & proper metrics — insight > accuracy chase"],
                ["CO5", "Certainty from tests + transparent explanations; report what is significant and what stays uncertain"],
            ],
            [1.2, 11.0],
            top=1.35,
        )

    content_slide(
        "What we learned in Foundations of Data Science",
        "CO1–CO5 as lived practice",
        s23,
        "If faculty asks what we learned in this subject — this table is our answer.",
        23,
        TOTAL,
    )

    # 24 Contribution matrix
    def s24(s):
        add_table(
            s,
            [
                ["Member", "One sentence to own in viva"],
                ["M1 Shasank", "I made churn a valid, time-separated label from cleaned retail events."],
                ["M2 [Name]", "I explored patterns and engineered observation-only RFM/behaviour features."],
                ["M3 [Name]", "I compared classifiers under a temporal split and froze the best evidenced model."],
                ["M4 [Name]", "I quantified feature–churn associations, explained predictions, and shipped the demo."],
            ],
            [2.8, 9.4],
            top=1.5,
        )

    content_slide(
        "Individual accountability",
        "Quick viva matrix",
        s24,
        "Each of us can point to files and numbers for that sentence.",
        24,
        TOTAL,
    )

    # 25 Closing
    def s25(s):
        # center message on closing slide
        box = s.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.3), Inches(3))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "Foundations first  →  models second  →  certainty always"
        set_run(r, size=28, bold=True, color=NAVY)

        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(28)
        r2 = p2.add_run()
        r2.text = "Thank you — questions welcome"
        set_run(r2, size=22, color=TEAL)

        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        p3.space_before = Pt(18)
        r3 = p3.add_run()
        r3.text = "Especially on: churn definition · leakage · EDA · statistical certainty"
        set_run(r3, size=16, color=MUTED)

    content_slide(
        "Closing",
        "Foundations of Data Science team",
        s25,
        "Thank you. Happy to take questions on churn definition, leakage control, EDA, or statistical certainty.",
        25,
        TOTAL,
    )

    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
